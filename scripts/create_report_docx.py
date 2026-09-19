import re
from pathlib import Path

from docx import Document


def markdown_to_docx(md_path: Path, docx_path: Path) -> None:
    doc = Document()
    lines = md_path.read_text(encoding="utf-8").splitlines()

    for raw_line in lines:
        line = raw_line.strip()

        if not line:
            doc.add_paragraph()
            continue

        if line.startswith("# "):
            doc.add_heading(line[2:], level=0)
            continue

        if line.startswith("## "):
            doc.add_heading(line[3:], level=1)
            continue

        if line.startswith("### "):
            doc.add_heading(line[4:], level=2)
            continue

        if line.startswith("#### "):
            doc.add_heading(line[5:], level=3)
            continue

        if line.startswith("- "):
            doc.add_paragraph(line[2:], style="List Bullet")
            continue

        paragraph = doc.add_paragraph()
        parts = re.split(r"(\*\*.*?\*\*)", line)
        for part in parts:
            if not part:
                continue
            run = paragraph.add_run(part[2:-2] if part.startswith("**") and part.endswith("**") else part)
            run.bold = part.startswith("**") and part.endswith("**")

    doc.save(docx_path)


if __name__ == "__main__":
    root = Path(__file__).resolve().parent.parent
    markdown_to_docx(
        root / "informe-salud-saladillo-marco-pba.md",
        root / "Informe-Salud-Saladillo-Marco-PBA.docx",
    )
