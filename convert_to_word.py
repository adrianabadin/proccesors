
import os
import re
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH

def markdown_to_docx(md_path, docx_path):
    if not os.path.exists(md_path):
        print(f"Error: No se encuentra el archivo {md_path}")
        return

    doc = Document()
    
    with open(md_path, 'r', encoding='utf-8') as f:
        lines = f.readlines()

    for line in lines:
        line = line.strip()
        
        # Saltos de línea vacíos
        if not line:
            doc.add_paragraph()
            continue
            
        # Títulos (Jerarquía)
        if line.startswith('# '):
            doc.add_heading(line[2:], level=0) # Título Principal
        elif line.startswith('## '):
            doc.add_heading(line[3:], level=1) # Título de Sección
        elif line.startswith('### '):
            doc.add_heading(line[4:], level=2) # Título de Ordenanza
        elif line.startswith('#### '):
            doc.add_heading(line[5:], level=3) # Subtítulos (Visto/Considerando)
            
        # Separadores
        elif line == '---':
            doc.add_page_break()
            
        # Listas
        elif line.startswith('- '):
            doc.add_paragraph(line[2:], style='List Bullet')
            
        # Texto en negrita (Articulado)
        elif line.startswith('**'):
            p = doc.add_paragraph()
            # Regex simple para detectar negritas básicas
            parts = re.split(r'(\*\*.*?\*\*)', line)
            for part in parts:
                if part.startswith('**') and part.endswith('**'):
                    p.add_run(part[2:-2]).bold = True
                else:
                    p.add_run(part)
        
        # Párrafos normales
        else:
            doc.add_paragraph(line)

    doc.save(docx_path)
    print(f"Documento Word generado exitosamente: {docx_path}")

if __name__ == "__main__":
    markdown_to_docx('ordenanzas-salud-completo.md', 'Compendio-Ordenanzas-Salud-Saladillo.docx')
